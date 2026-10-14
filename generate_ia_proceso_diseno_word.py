import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_doc():
    doc = docx.Document()
    
    # Page Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)
        sec.header.is_linked_to_previous = False
        
    navy = RGBColor(10, 28, 52)         # #0A1C34
    coral = RGBColor(207, 82, 36)       # #CF5224
    gray_dark = RGBColor(31, 41, 55)    # #1F2937
    gray_muted = RGBColor(100, 116, 139)# #64748B
    
    # Title Header Box (Table)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.7)
    c = tbl.cell(0, 0)
    set_cell_shading(c, '0A1C34')
    set_cell_margins(c, top=180, bottom=180, left=220, right=220)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r_tag = p.add_run('MÓDULO ESTRATÉGICO ACADÉMICO | SEGUROS DEL ESTADO & MANTO 2.0')
    r_tag.font.name = 'Arial'
    r_tag.font.size = Pt(8.5)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(207, 82, 36)
    
    p_t = c.add_paragraph()
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run('IA y el Proceso de Diseño en el Sector Asegurador')
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(255, 255, 255)
    
    p_sub = c.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(0)
    r_sub = p_sub.add_run('Respuestas Ejecutivas — Límite de 300 Palabras (Conteo: 289 palabras con títulos / 267 cuerpo)')
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(203, 213, 225)
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(8)
    
    # Helper to add section
    def add_question(num_str, title_str, content_paragraphs):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        r_num = h.add_run(num_str + ' ')
        r_num.font.name = 'Arial'
        r_num.font.size = Pt(11)
        r_num.font.bold = True
        r_num.font.color.rgb = coral
        
        r_title = h.add_run(title_str)
        r_title.font.name = 'Arial'
        r_title.font.size = Pt(11)
        r_title.font.bold = True
        r_title.font.color.rgb = navy
        
        for cp in content_paragraphs:
            p_c = doc.add_paragraph()
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(3)
            p_c.paragraph_format.line_spacing = 1.15
            for run_item in cp:
                text, bold, italic, bullet = run_item
                if bullet:
                    p_c.paragraph_format.left_indent = Inches(0.2)
                r = p_c.add_run(text)
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.font.bold = bold
                r.font.italic = italic
                r.font.color.rgb = gray_dark

    # Q1
    q1_content = [
        [('El sector asegurador colombiano avanza hacia la digitalización. Seguros del Estado, aseguradora líder multirramo, impulsa desde su Gerencia de Productividad y Transformación de Procesos la optimización operativa mediante rediseño de procesos, automatización inteligente y agilidad para maximizar su competitividad.', False, False, False)]
    ]
    add_question('1.', 'Contexto: Sector Asegurador y Transformación de Procesos', q1_content)
    
    # Q2
    q2_content = [
        [('Debe operar bajo un modelo ágil, customer-centric y modular para comprimir el time-to-market:', False, False, False)],
        [('• Discovery y Testeo Continuo: ', True, False, True), ('Entrevistas con usuarios reales para resolver dolores genuinos antes de estructurar pólizas.', False, False, False)],
        [('• Co-creación Interdisciplinaria: ', True, False, True), ('Mesas ágiles integrando simultáneamente actuaría, suscripción, legal, tecnología y UX.', False, False, False)],
        [('• Arquitectura Composable (API-First): ', True, False, True), ('Coberturas atómicas desacopladas como microservicios parametrizables según el riesgo y presupuesto del asegurado.', False, False, False)],
        [('• Experimentación Lean: ', True, False, True), ('MVPs iterativos validados en campo antes del escalamiento masivo.', False, False, False)]
    ]
    add_question('2.', 'Proceso de Diseño: Ágil, Customer-Centric y Modular', q2_content)
    
    # Q3
    q3_content = [
        [('MANTO es el spin-off insurtech de Seguros del Estado para digitalizar integralmente la suscripción y emisión de seguros (iniciando en mascotas). Opera bajo una arquitectura híbrida donde la IA percibe, conversa y estructura, mientras el core actuarial decide:', False, False, False)],
        [('• Nuquí (LLMs de Triaje): ', True, False, True), ('Agente conversacional empático (Web/WhatsApp) que interactúa en lenguaje natural, perfila al asegurado y traduce el diálogo en datos estructurados (JSON) de suscripción, eliminando formularios engorrosos.', False, False, False)],
        [('• Google Gemini Multimodal: ', True, False, True), ('Visión por computador en tiempo real que analiza la fotografía de la mascota para clasificar biológicamente razas y fenotipos, y detectar obligatoriamente razas de manejo especial bajo Ley 1801 de 2016, impidiendo falsas autodeclaraciones.', False, False, False)],
        [('• Sué (Visión + OCR DIAN): ', True, False, True), ('Extracción en facturas veterinarias y validación directa de CUFE ante la DIAN, automatizando reembolsos y blindando contra fraude tributario.', False, False, False)],
        [('• Motor Actuarial Determinista: ', True, False, True), ('Tarificación 100% matemática en servidor, garantizando cero alucinaciones y total auditabilidad.', False, False, False)],
        [('Impacto: ', True, False, False), ('Emisión inmediata Straight-Through Processing (STP) en minutos, suscripción sin sesgos, blindaje antifraude y experiencia sin fricción.', False, False, False)]
    ]
    add_question('3.', 'Integración con IA: MANTO como Spin-off de Seguros del Estado', q3_content)
    
    # Callout Footer Note Box
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    tbl_note = doc.add_table(rows=1, cols=1)
    tbl_note.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_note.autofit = False
    tbl_note.columns[0].width = Inches(6.7)
    cn = tbl_note.cell(0, 0)
    set_cell_shading(cn, 'F1F5F9')
    set_cell_margins(cn, top=80, bottom=80, left=130, right=130)
    pn = cn.paragraphs[0]
    pn.paragraph_format.space_before = Pt(0)
    pn.paragraph_format.space_after = Pt(0)
    rn1 = pn.add_run('Gobernanza y Cumplimiento: ')
    rn1.font.name = 'Arial'
    rn1.font.size = Pt(8.5)
    rn1.font.bold = True
    rn1.font.color.rgb = navy
    rn2 = pn.add_run('Iniciativa Spin-off de Seguros del Estado bajo directrices SFC, Ley 1801 de 2016, Ley 1581 de 2012 y suscripción actuarial determinista. Conteo estricto: 267 palabras en cuerpo / 289 con títulos (Límite máximo: 300 palabras).')
    rn2.font.name = 'Arial'
    rn2.font.size = Pt(8.5)
    rn2.font.color.rgb = gray_muted
    
    out_dir = r'e:\Manto 2.0'
    file_path = os.path.join(out_dir, 'IA_y_el_proceso_de_diseno_Seguros_del_Estado_Manto.docx')
    doc.save(file_path)
    docs_path = os.path.join(out_dir, 'docs', 'IA_y_el_proceso_de_diseno_Seguros_del_Estado_Manto.docx')
    doc.save(docs_path)
    print(f'Document successfully created at {file_path} and {docs_path}')

if __name__ == '__main__':
    create_doc()
