import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    navy = RGBColor(10, 28, 52)      # #0A1C34
    orange = RGBColor(207, 82, 36)   # #CF5224
    gray_muted = RGBColor(74, 90, 111) # #4A5A6F

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("PLAN MAESTRO DE PRODUCTO Y ARQUITECTURA\nMANTO 2.0")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = navy

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Plataforma Insurtech Digital Multi-Ramo para Colombia y Latinoamérica\nVersión 2.0.0 — Documento de Gobernanza y Especificación Técnica")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = gray_muted

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 1. Resumen Ejecutivo
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Resumen Ejecutivo y Visión de Producto")
    r1.font.color.rgb = navy

    p = doc.add_paragraph(
        "Manto 2.0 es la plataforma de seguros digitales (Insurtech) diseñada para transformar la experiencia de aseguramiento en Colombia y la región. "
        "A diferencia de las aseguradoras tradicionales con pólizas rígidas en papel y tiempos de radicación de semanas, Manto ofrece:"
    )
    doc.add_paragraph("• Cotización y Personalización en < 60 segundos con asistencia de Nuquí (IA multimodal + motor determinista).", style='List Bullet')
    doc.add_paragraph("• Tarificación dinámica a la medida (Tailored Baseline) eliminando los paquetes rígidos de Bronce/Plata/Oro.", style='List Bullet')
    doc.add_paragraph("• Radicación y Reembolsos Veterinarios Automatizados con Sué (Validación CUFE DIAN + Aprobación STP en minutos para montos ≤ $500.000 COP).", style='List Bullet')
    doc.add_paragraph("• Cumplimiento legal estricto: Ley 1801 de 2016 (RCE 50 SMLMV para razas especiales), Habeas Data (Ley 1581) y dinero nominal exacto (ColombianPeso).", style='List Bullet')

    # 2. Arquitectura Tecnológica
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("2. Arquitectura Tecnológica y Stack")
    r1.font.color.rgb = navy

    p = doc.add_paragraph(
        "La plataforma está construida bajo el patrón de Monolito Modular con principios de Domain-Driven Design (DDD), garantizando alta cohesión, "
        "bajo acoplamiento y facilidad de evolución hacia microservicios si la escala lo requiere:"
    )

    table_stack = doc.add_table(rows=5, cols=2)
    table_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    stack_data = [
        ("Capa", "Tecnología y Justificación"),
        ("Backend", "NestJS 10 (TypeScript 5, Node.js 20+). Monolito modular con DDD, CQRS ligero, Event Emitter y validación por DTOs class-validator."),
        ("Frontend", "Next.js 14 (App Router, React 18, Tailwind CSS, @fontsource/poppins). Mobile-first, diseño Sanctuary, micro-animaciones fluidas."),
        ("Base de Datos", "PostgreSQL 16 sin ORM (driver pg nativo). Migraciones transaccionales, UUIDv7, dinero como BigInt/cadena exacta y datos personales en columnas."),
        ("Inteligencia Artificial", "Gemini 2.0 Flash Vision con fallback determinista de 3s. La IA asiste y clasifica, pero NUNCA es autoridad de suscripción ni calcula primas.")
    ]
    for row_idx, (col1, col2) in enumerate(stack_data):
        row = table_stack.rows[row_idx]
        row.cells[0].text = col1
        row.cells[1].text = col2
        if row_idx == 0:
            set_cell_shading(row.cells[0], "0A1C34")
            set_cell_shading(row.cells[1], "0A1C34")
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
        else:
            bg = "F4F6F8" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_shading(row.cells[0], bg)
            set_cell_shading(row.cells[1], bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. Dominios del Sistema
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("3. Los Dominios Core de la Plataforma")
    r1.font.color.rgb = navy

    p = doc.add_paragraph("El sistema desacopla sus responsabilidades en 8 módulos core altamente cohesivos:")

    domains = [
        ("Pricing Engine (Tarificación)", "Motor determinista que cruza especie, raza, contextura, edad, sexo, ciudad y deducible. Gestiona el catálogo actuarial y amparos."),
        ("Quotes (Cotizaciones)", "Generación de cotizaciones inmutables con UUIDv7 y vigencia de 15 días calendario. Control de elegibilidad."),
        ("Policies (Pólizas)", "Emisión, activación, estado de vigencia y ciclo de vida de la póliza contratada."),
        ("Claims & Siniestros (Sué ⚡)", "Radicación FNOL, máquina de estados de siniestros, motor STP (≤ $500.000 COP) y auditoría médica veterinaria."),
        ("Reimbursements (Pagos)", "Dispersión de fondos a través de Bre-B (llave celular/cédula) y pasarelas de pago."),
        ("AI & Conversational (Nuquí 🐾)", "Visión por computador para detección biométrica de mascotas y orquestación del chat de cotización sin fricción."),
        ("Eligibility (Suscripción)", "Reglas de nuevo ingreso (2 a 144 meses) y validaciones de asegurabilidad."),
        ("Identity & Security", "Habeas Data (derecho de supresión Ley 1581), control de acceso y prevención de fraude.")
    ]

    t_dom = doc.add_table(rows=len(domains)+1, cols=2)
    t_dom.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_dom.rows[0].cells[0].text = "Módulo / Dominio"
    t_dom.rows[0].cells[1].text = "Responsabilidad Clave"
    set_cell_shading(t_dom.rows[0].cells[0], "0A1C34")
    set_cell_shading(t_dom.rows[0].cells[1], "0A1C34")
    for c in t_dom.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    for i, (d_name, d_desc) in enumerate(domains, start=1):
        row = t_dom.rows[i]
        row.cells[0].text = d_name
        row.cells[1].text = d_desc
        bg = "F4F6F8" if i % 2 == 1 else "FFFFFF"
        set_cell_shading(row.cells[0], bg)
        set_cell_shading(row.cells[1], bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Estructura de Tarificación (Pricing)
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("4. Estructura de Cobro de Prima y Motor Actuarial")
    r1.font.color.rgb = navy

    doc.add_paragraph(
        "Fórmula Maestra:\n"
        "Prima Base Mensual = Tasa Pura ($30.000) × F_especie × F_raza × F_tamaño × F_sexo × F_edad × F_ciudad × F_deducible × (1 - D_multimascota)\n"
        "Prima Total Mensual = Prima Base Mensual + ∑ Primas de Add-ons Activos\n"
        "Prima Total Anual = Prima Total Mensual × 11 (1 mes gratis por pago anticipado)"
    )

    t_addons = doc.add_table(rows=7, cols=4)
    t_addons.alignment = WD_TABLE_ALIGNMENT.CENTER
    addons_data = [
        ("Add-on / Cobertura Opcional", "Cuota Mensual", "Límite Anual", "Días Activación"),
        ("Enfermedades Comunes e Infecciosas", "$24.900 COP", "$5.000.000 COP", "30 días"),
        ("Medicamentos Recetados Ambulatorios", "$12.900 COP", "$600.000 COP", "30 días"),
        ("Responsabilidad Civil (Ley 1801)", "$15.000 COP", "$87.545.250 (50 SMLMV)", "Inmediata (0 d)"),
        ("Asistencia Exequial y Cremación", "$9.900 COP", "$1.200.000 COP", "30 días"),
        ("Telemedicina Veterinaria 24/7", "$6.900 COP", "Ilimitado", "Inmediata (0 d)"),
        ("Guardería por Hospitalización Tomador", "$4.900 COP", "$800.000 COP", "15 días")
    ]
    for row_idx, row_values in enumerate(addons_data):
        row = t_addons.rows[row_idx]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].text = val
            if row_idx == 0:
                set_cell_shading(row.cells[col_idx], "0A1C34")
                for p in row.cells[col_idx].paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F4F6F8" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_shading(row.cells[col_idx], bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. Roadmap y Estado
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("5. Estado del Desarrollo y Roadmap")
    r1.font.color.rgb = navy

    doc.add_paragraph("Fase 1: MVP Core Mascotas (Completada):")
    doc.add_paragraph("✓ Motor de tarificación determinista en NestJS con nominales BigInt.", style='List Bullet')
    doc.add_paragraph("✓ Flujo conversacional de Nuquí en Next.js con carga de fotos y selección de amparos.", style='List Bullet')
    doc.add_paragraph("✓ Migraciones PostgreSQL y persistencia con trazabilidad de RCE Ley 1801.", style='List Bullet')

    doc.add_paragraph("Fase 2: Checkout, Segunda Mascota y Siniestros (En curso):")
    doc.add_paragraph("• Flujo de alta de segunda mascota con 5% de descuento.", style='List Bullet')
    doc.add_paragraph("• Checkout y pasarela de pago (Wompi / PSE / Tarjeta / Bre-B).", style='List Bullet')
    doc.add_paragraph("• Portal de reembolsos con Sué (validación DIAN CUFE + STP).", style='List Bullet')

    # Output file
    doc.save("e:/Manto 2.0/PLAN_MAESTRO_MANTO_2.0.docx")
    print("Document successfully created: e:/Manto 2.0/PLAN_MAESTRO_MANTO_2.0.docx")

if __name__ == '__main__':
    create_document()
